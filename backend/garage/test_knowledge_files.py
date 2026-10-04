import io
import json
from unittest.mock import MagicMock, patch

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from rest_framework.test import APIClient
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from .knowledge_files import KnowledgeFileError, extract_file
from .models import KnowledgeDocument


class ExtractionTests(SimpleTestCase):
    def file(self, name, data):
        return SimpleUploadedFile(name, data)

    def test_text_csv_and_json(self):
        cases = [
            ("manual.txt", "สเปกเครื่องยนต์".encode(), "สเปกเครื่องยนต์"),
            ("manual.md", b"# CB500\nEngine 471 cc", "Engine 471 cc"),
            ("spec.csv", b"model,cc\nCB500,471\n", "CB500 | 471"),
            ("spec.json", json.dumps({"model": "CB500", "cc": 471}).encode(), "CB500"),
        ]
        for name, data, expected in cases:
            with self.subTest(name=name):
                result = extract_file(self.file(name, data))
                self.assertIn(expected, result["extracted_text"])
                self.assertFalse(result["sections"][0]["reviewed"])

    def test_docx_and_xlsx(self):
        document = Document()
        document.add_paragraph("CB500 engine 471 cc")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Model"
        table.cell(0, 1).text = "CB500"
        output = io.BytesIO()
        document.save(output)
        text = extract_file(self.file("manual.docx", output.getvalue()))["extracted_text"]
        self.assertIn("CB500 engine 471 cc", text)
        self.assertIn("Model | CB500", text)

        workbook = Workbook()
        workbook.active.append(["Model", "CC"])
        workbook.active.append(["CB500", 471])
        output = io.BytesIO()
        workbook.save(output)
        text = extract_file(self.file("spec.xlsx", output.getvalue()))["extracted_text"]
        self.assertIn("CB500 | 471", text)

    def test_text_pdf_keeps_page_locator(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
            DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 72 720 Td (CB500 engine 471 cc) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        output = io.BytesIO()
        writer.write(output)
        result = extract_file(self.file("manual.pdf", output.getvalue()))
        self.assertIn("CB500 engine 471 cc", result["extracted_text"])
        self.assertIn("หน้า 1", result["sections"][0]["locator"])

    def test_rejects_wrong_type_and_empty_pdf(self):
        with self.assertRaises(KnowledgeFileError):
            extract_file(self.file("script.exe", b"payload"))
        with self.assertRaises(KnowledgeFileError):
            extract_file(self.file("manual.pdf", b"not a pdf"))
        with self.assertRaises(KnowledgeFileError):
            extract_file(self.file("broken.pdf", b"%PDF-corrupt"))


class DocumentApiTests(TestCase):
    def setUp(self):
        cache.clear()
        users = get_user_model()
        self.admin = users.objects.create_superuser(username="document_admin", password="test")
        self.customer = users.objects.create_user(username="document_customer", password="test")
        self.client = APIClient()
        self.client.force_authenticate(self.admin)

    def upload(self):
        return self.client.post("/api/admin/knowledge-documents/", {
            "title": "CB500 specifications", "brand": "Honda", "model": "CB500",
            "source_url": "https://example.test/spec", "file": SimpleUploadedFile(
                "cb500.txt", "Engine capacity 471 cc".encode()),
        }, format="multipart")

    def test_admin_upload_review_embed_edit_and_archive(self):
        created = self.upload()
        self.assertEqual(created.status_code, 201, created.data)
        pk = created.data["id"]
        detail_url = f"/api/admin/knowledge-documents/{pk}/"
        detail = self.client.get(detail_url)
        self.assertIn("Engine capacity 471 cc", detail.data["extracted_text"])
        self.assertEqual(detail.data["section_count"], 1)
        section_url = detail_url + "sections/0/"
        self.assertEqual(self.client.post(section_url + "embed/").status_code, 400)
        self.client.force_authenticate(self.customer)
        self.assertEqual(self.client.get(detail_url).status_code, 403)
        self.assertEqual(self.client.post("/api/admin/knowledge-documents/", {
            "title": "Wrong", "file": SimpleUploadedFile("a.txt", b"text"),
        }, format="multipart").status_code, 403)
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.client.patch(section_url, {"reviewed": True},
                                           format="json").status_code, 200)
        with patch("garage.admin_documents.embedding_key", return_value="test"), \
             patch("garage.admin_documents.embed_content", return_value=[0.1] * 3072), \
             patch("garage.admin_documents.vector_connection") as connect:
            connect.return_value.__enter__.return_value.cursor.return_value.__enter__.return_value = MagicMock()
            embedded = self.client.post(section_url + "embed/")
            self.assertEqual(embedded.status_code, 200, embedded.data)
            self.assertEqual(embedded.data["embedding_status"], "ready")
            edited = self.client.patch(section_url, {"content": "Corrected 471 cc"}, format="json")
            self.assertEqual(edited.status_code, 200, edited.data)
            self.assertFalse(edited.data["reviewed"])
            self.assertEqual(edited.data["embedding_status"], "pending")
            self.assertTrue(any("DELETE FROM the_x_manual_vectors" in call.args[0]
                for call in connect.return_value.__enter__.return_value.cursor.return_value.__enter__.return_value.execute.call_args_list))
            self.assertEqual(self.client.delete(detail_url).status_code, 200)
        self.assertEqual(self.client.get(detail_url).status_code, 404)
        self.assertIsNotNone(KnowledgeDocument.objects.get(pk=pk).archived_at)

    def test_upload_rejects_unsupported_or_empty_content(self):
        response = self.client.post("/api/admin/knowledge-documents/", {
            "title": "Invalid", "file": SimpleUploadedFile("bad.pdf", b"wrong"),
        }, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(KnowledgeDocument.objects.exists())
